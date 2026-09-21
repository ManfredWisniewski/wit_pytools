import csv
import importlib
import json
import os
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

# Allow importing wit_pytools when running tests directly
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../..')))

from wit_pytools.documenttools import (
    anonymize_xlsx,
    create_xlsx_mapping,
    document_find_regex,
    identify_xlsx_strings,
    pdf_to_markdown,
    pdf_to_markdown_text,
)

TEST_DOC = Path(__file__).parent / "documenttools" / "testdocument.pdf"


def test_document_find_regex_literal_real_document():
    results = document_find_regex(TEST_DOC, "manfred@mustermann.de", context_chars=10)

    assert len(results) >= 1
    assert results[0]["match"].lower() == "manfred@mustermann.de"


def test_document_find_regex_regex_real_document():
    pattern = r"[A-Za-z]+@mustermann\.de"
    results = document_find_regex(TEST_DOC, pattern, regex=True)

    assert any(res["match"].lower().endswith("@mustermann.de") for res in results)


def test_document_find_regex_requires_pdfplumber(monkeypatch):
    monkeypatch.setattr("wit_pytools.documenttools.pdfplumber", None)

    with pytest.raises(RuntimeError):
        document_find_regex(TEST_DOC, "anything")


def test_xlsx_anonymization_workflow_preserves_formulas(tmp_path):
    source_path = tmp_path / "source.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet["A1"] = "Alice Example"
    worksheet["A2"] = "Alice Example"
    worksheet["A3"] = "alice@example.com"
    worksheet["A4"] = 42
    worksheet["A5"] = "=SUM(A4,1)"
    worksheet["A6"] = "31.03.2020"
    workbook.save(source_path)

    candidates_path = identify_xlsx_strings(source_path)
    with candidates_path.open(encoding="utf-8", newline="") as file_handle:
        candidate_rows = list(csv.DictReader(file_handle))

    assert [row["original_value"] for row in candidate_rows] == [
        "Alice Example",
        "alice@example.com",
    ]
    assert "31.03.2020" not in [
        row["original_value"] for row in candidate_rows
    ]
    assert candidate_rows[0]["occurrences"] == "2"
    assert candidate_rows[0]["cell"] == "A1; A2"

    candidate_rows[0]["replacement_value"] = "Person-1"
    candidate_rows[1]["replacement_value"] = "person-1@example.invalid"
    with candidates_path.open("w", encoding="utf-8", newline="") as file_handle:
        writer = csv.DictWriter(file_handle, fieldnames=candidate_rows[0].keys())
        writer.writeheader()
        writer.writerows(candidate_rows)

    mapping_path = create_xlsx_mapping(candidates_path)
    output_path = anonymize_xlsx(source_path, mapping_path)

    anonymized = openpyxl.load_workbook(output_path, data_only=False)
    anonymized_sheet = anonymized["Data"]
    assert anonymized_sheet["A1"].value == "Person-1"
    assert anonymized_sheet["A2"].value == "Person-1"
    assert anonymized_sheet["A3"].value == "person-1@example.invalid"
    assert anonymized_sheet["A4"].value == 42
    assert anonymized_sheet["A5"].value == "=SUM(A4,1)"
    assert anonymized_sheet["A6"].value == "31.03.2020"
    anonymized.close()


def test_create_xlsx_mapping_rejects_invalid_rows(tmp_path):
    candidate_path = tmp_path / "source_candidates.csv"
    with candidate_path.open("w", encoding="utf-8", newline="") as file_handle:
        writer = csv.DictWriter(
            file_handle,
            fieldnames=[
                "original_value",
                "replacement_value",
                "value_type",
                "worksheet",
                "cell",
                "occurrences",
            ],
        )
        writer.writeheader()
        writer.writerow(
            {
                "original_value": "same",
                "replacement_value": "one",
                "value_type": "string",
                "worksheet": "Sheet",
                "cell": "A1",
                "occurrences": "1",
            }
        )
        writer.writerow(
            {
                "original_value": "same",
                "replacement_value": "two",
                "value_type": "string",
                "worksheet": "Sheet",
                "cell": "A2",
                "occurrences": "1",
            }
        )

    with pytest.raises(ValueError, match="Duplicate original_value"):
        create_xlsx_mapping(candidate_path)


def test_create_xlsx_mapping_groups_contained_values(tmp_path):
    candidate_path = tmp_path / "source_candidates.csv"
    with candidate_path.open("w", encoding="utf-8", newline="") as file_handle:
        writer = csv.DictWriter(file_handle, fieldnames=[
            "original_value",
            "replacement_value",
            "value_type",
            "worksheet",
            "cell",
            "occurrences",
        ])
        writer.writeheader()
        for value, replacement, value_type in [
            ("Lia", "value-lia", "string"),
            ("Barfuss", "value-bar", "string"),
            ("Lia Barfuss", "Person-54b", "name"),
            ("Herr Barfuss", "value-herr", "name"),
        ]:
            writer.writerow({
                "original_value": value,
                "replacement_value": replacement,
                "value_type": value_type,
                "worksheet": "Sheet",
                "cell": "A1",
                "occurrences": "1",
            })

    mapping_path = create_xlsx_mapping(
        candidate_path,
        group_contained_values=True,
    )
    with mapping_path.open(encoding="utf-8", newline="") as file_handle:
        reader = csv.DictReader(file_handle)
        mapping_rows = list(reader)

    assert reader.fieldnames == [
        "replacement_value",
        "original_value",
        "value_type",
    ]
    assert len(mapping_rows) == 1
    assert mapping_rows[0]["original_value"] == (
        "Lia;Barfuss;Lia Barfuss;Herr Barfuss"
    )
    assert {row["replacement_value"] for row in mapping_rows} == {"Person-54b"}


def test_anonymize_xlsx_preserves_personenkonten_substring_matches(tmp_path):
    source_path = tmp_path / "source.xlsx"
    mapping_path = tmp_path / "source_mapping.csv"
    output_path = tmp_path / "source_anonymized.xlsx"

    workbook = openpyxl.Workbook()
    personenkonten = workbook.active
    personenkonten.title = "Personenkonten"
    personenkonten["F3"] = "Barfuss"
    personenkonten["AM3"] = "Prüfung:"
    kontoauszuege = workbook.create_sheet("Kontoauszüge")
    kontoauszuege["F1"] = "Klassenkasse 6A Lia Barfuss RINP Dauerauftrag"
    workbook.save(source_path)

    with mapping_path.open("w", encoding="utf-8", newline="") as file_handle:
        writer = csv.DictWriter(
            file_handle,
            fieldnames=["original_value", "replacement_value", "value_type"],
        )
        writer.writeheader()
        writer.writerows(
            [
                {
                    "original_value": "Barfuss",
                    "replacement_value": "value-48a",
                    "value_type": "string",
                },
                {
                    "original_value": "Klassenkasse 6A Lia Barfuss RINP Dauerauftrag",
                    "replacement_value": "value-48a",
                    "value_type": "string",
                },
            ]
        )

    anonymize_xlsx(source_path, mapping_path, output_path)

    anonymized = openpyxl.load_workbook(output_path, data_only=False)
    assert anonymized["Personenkonten"]["F3"].value == "value-48a"
    assert anonymized["Personenkonten"]["AM3"].value == "Prüfung:"
    assert anonymized["Kontoauszüge"]["F1"].value == "value-48a"
    anonymized.close()


# --- PDF to Markdown -------------------------------------------------------

pdf2md = importlib.import_module("wit_pytools.documenttools.pdf2md")


@pytest.fixture
def pdf_copy(tmp_path):
    target = tmp_path / "testdocument.pdf"
    shutil.copy(TEST_DOC, target)
    return target


@pytest.fixture
def no_cost_prompt(monkeypatch):
    monkeypatch.setattr(pdf2md, "estimate_cost", lambda model, pages, api_key=None: 0.0)
    monkeypatch.setattr(pdf2md.time, "sleep", lambda seconds: None)


def _fake_chat(responses):
    calls = []

    def chat(messages, model, *, system=None, api_key=None):
        calls.append({"messages": messages, "model": model, "system": system})
        response = responses.pop(0)
        if isinstance(response, Exception):
            raise response
        return response

    chat.calls = calls
    return chat


def test_pdf_to_markdown_vision_writes_output_and_sidecar(pdf_copy, monkeypatch, no_cost_prompt):
    fake = _fake_chat(["```markdown\n# Seite 1\n```", "Seite 2 Text"])
    monkeypatch.setattr(pdf2md, "chat", fake)

    output = pdf_to_markdown(pdf_copy, model="test/vision", language="de")

    assert output == pdf_copy.with_suffix(".md")
    assert output.read_text(encoding="utf-8") == "# Seite 1\n\n---\n\nSeite 2 Text\n"
    assert len(fake.calls) == 2
    assert fake.calls[0]["model"] == "test/vision"
    prompt_text = fake.calls[0]["messages"][0]["content"][0]["text"]
    assert "[unleserlich]" in prompt_text and "{illegible}" not in prompt_text
    assert fake.calls[0]["messages"][0]["content"][1]["type"] == "image_url"

    sidecar = json.loads(pdf_copy.with_name("testdocument_pdf2md.json").read_text(encoding="utf-8"))
    assert sidecar["mode"] == "vision"
    assert sidecar["pages"] == [1, 2]
    assert [entry["status"] for entry in sidecar["page_status"]] == ["ok", "ok"]
    assert not list(pdf_copy.parent.glob("page_*.png"))


def test_pdf_to_markdown_supports_separate_sidecar(pdf_copy, tmp_path):
    sidecar = tmp_path / "originals" / "testdocument_pdf2md.json"

    output = pdf_to_markdown(
        pdf_copy,
        mode="text",
        sidecar_path=sidecar,
    )

    assert output == pdf_copy.with_suffix(".md")
    assert sidecar.is_file()
    assert not pdf_copy.with_name("testdocument_pdf2md.json").exists()


def test_pdf_to_markdown_refuses_overwrite(pdf_copy, monkeypatch, no_cost_prompt):
    monkeypatch.setattr(pdf2md, "chat", _fake_chat(["a", "b", "c", "d"]))
    pdf_to_markdown(pdf_copy, model="m")

    with pytest.raises(FileExistsError):
        pdf_to_markdown(pdf_copy, model="m")
    pdf_to_markdown(pdf_copy, model="m", overwrite=True)


def test_pdf_to_markdown_retries_then_marks_failed_page(pdf_copy, monkeypatch, no_cost_prompt):
    fake = _fake_chat(["", RuntimeError("OpenRouter 502: bad gateway"), "second try", "page two"])
    monkeypatch.setattr(pdf2md, "chat", fake)

    text = pdf_to_markdown_text(pdf_copy, model="m", retry_times=3)

    assert text == "second try\n\n---\n\npage two"
    assert len(fake.calls) == 4


def test_pdf_to_markdown_failed_page_raises_unless_continue(pdf_copy, monkeypatch, no_cost_prompt):
    errors = [RuntimeError("boom")] * 2
    monkeypatch.setattr(pdf2md, "chat", _fake_chat(errors + ["ok"]))
    with pytest.raises(RuntimeError, match="page\\(s\\): 1"):
        pdf_to_markdown_text(pdf_copy, model="m", retry_times=2)

    monkeypatch.setattr(pdf2md, "chat", _fake_chat([RuntimeError("boom")] * 2 + ["ok"]))
    text = pdf_to_markdown_text(pdf_copy, model="m", retry_times=2, continue_on_error=True)
    assert text.startswith("> **Page 1: conversion failed** — boom")
    assert text.endswith("ok")


def test_pdf_to_markdown_text_mode_uses_text_layer_without_api(pdf_copy, monkeypatch):
    monkeypatch.setattr(pdf2md, "chat", _fake_chat([]))

    text = pdf_to_markdown_text(pdf_copy, mode="text", end_page=1)

    assert "Unternehmensanmeldung" in text
    assert "---" not in text


def test_pdf_to_markdown_keep_pages_and_page_range(pdf_copy, monkeypatch, no_cost_prompt):
    monkeypatch.setattr(pdf2md, "chat", _fake_chat(["only page two"]))

    pdf_to_markdown(pdf_copy, model="m", start_page=2, end_page=2, keep_pages=True)

    pages_dir = pdf_copy.with_name("testdocument_pages")
    assert (pages_dir / "page_0002.png").is_file()
    assert (pages_dir / "page_0002.md").read_text(encoding="utf-8") == "only page two"
    assert not (pages_dir / "page_0001.png").exists()


def test_pdf_to_markdown_validation(pdf_copy, monkeypatch):
    with pytest.raises(ValueError):
        pdf_to_markdown_text(pdf_copy, mode="text", start_page=3)
    with pytest.raises(ValueError):
        pdf_to_markdown_text(pdf_copy, mode="ocr")
    with pytest.raises(ValueError, match="Unsupported language"):
        pdf_to_markdown_text(pdf_copy, mode="text", language="fr")
    monkeypatch.delenv("OPENROUTER_PDF_MODEL", raising=False)
    monkeypatch.delenv("OPENROUTER_MODEL", raising=False)
    with pytest.raises(RuntimeError, match="OPENROUTER_PDF_MODEL"):
        pdf_to_markdown_text(pdf_copy)


def test_pdf_to_markdown_image_only_pdf(tmp_path, monkeypatch, no_cost_prompt):
    from PIL import Image

    scan = tmp_path / "scan.pdf"
    Image.new("RGB", (200, 300), "white").save(scan, "PDF")

    text = pdf_to_markdown_text(scan, mode="text", language="de")
    assert text == "> **Seite 1: keine Textebene** — für diese Seite den Modus vision verwenden"

    monkeypatch.setattr(pdf2md, "chat", _fake_chat(["transcribed scan"]))
    assert pdf_to_markdown_text(scan, model="m") == "transcribed scan"


def test_pdf_to_markdown_cost_confirmation_aborts(pdf_copy, monkeypatch):
    monkeypatch.setattr(pdf2md, "estimate_cost", lambda model, pages, api_key=None: 5.0)
    monkeypatch.setattr("builtins.input", lambda prompt: "n")
    monkeypatch.setattr(pdf2md, "chat", _fake_chat([]))

    with pytest.raises(RuntimeError, match="Aborted"):
        pdf_to_markdown_text(pdf_copy, model="m", max_cost=1.0)

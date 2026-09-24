import argparse
import tempfile
from pathlib import Path

import openpyxl

from wit_pytools.anonymization import mapping_path_rows
from wit_pytools.documenttools import anonymize_xlsx


def main() -> int:
    parser = argparse.ArgumentParser(
        description="List cells that changed although their value is not in the mapping."
    )
    parser.add_argument("source", type=Path)
    parser.add_argument("mapping", type=Path)
    args = parser.parse_args()

    mapping = mapping_path_rows(args.mapping)
    with tempfile.TemporaryDirectory() as tmp:
        output = anonymize_xlsx(args.source, args.mapping, Path(tmp) / "out.xlsx")
        src = openpyxl.load_workbook(args.source, data_only=False)
        dst = openpyxl.load_workbook(output, data_only=False)
        changed = 0
        for ws in src.worksheets:
            ws_out = dst[ws.title]
            for row in ws.iter_rows():
                for cell in row:
                    before = cell.value
                    after = ws_out[cell.coordinate].value
                    if before == after or not isinstance(before, str):
                        continue
                    changed += 1
                    if before not in mapping:
                        print(f"{ws.title}!{cell.coordinate}: {before!r} -> {after!r}")
        print(f"\nTotal changed string cells: {changed}")
        src.close()
        dst.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

"""Document-specific search, anonymization adapters, and PDF helpers."""

import csv
import os
import re
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from eliot import log_message

from wit_pytools.anonymization import (
    CANDIDATE_COLUMNS,
    CandidateCollector,
    detect_presidio_candidates,
    detect_text_candidates,
    is_date_string,
    load_name_catalog,
    mapping_matches_parts,
    mapping_path_rows,
    replace_related_values,
    replace_text_parts,
    related_mapping_values,
    write_csv,
)

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

try:
    import pdfplumber  # type: ignore
except ImportError:  # pragma: no cover
    pdfplumber = None

_LAZY_EXPORTS = {"pdf_to_markdown", "pdf_to_markdown_text"}


def __getattr__(name: str):
    if name in _LAZY_EXPORTS:
        import importlib

        return getattr(importlib.import_module(f"{__name__}.pdf2md"), name)
    raise AttributeError(f"module {__name__!r} has no attribute {name!r}")


def document_find_regex(
    file_path: Path | str,
    query: str,
    *,
    regex: bool = False,
    context_chars: int = 40,
    max_pages: Optional[int] = None,
    flags: int = re.IGNORECASE,
) -> List[Dict[str, Any]]:
    """Search a PDF for a literal string or regex pattern and return contexts."""
    if not pdfplumber:
        raise RuntimeError(
            "pdfplumber is required for document_find_regex. "
            "Install pdfplumber to use this function."
        )

    pattern = re.compile(query if regex else re.escape(query), flags)
    pdf_path = Path(file_path)
    results: List[Dict[str, Any]] = []
    try:
        with pdfplumber.open(str(pdf_path)) as pdf:
            pages = pdf.pages if max_pages is None else pdf.pages[:max_pages]
            for page_index, page in enumerate(pages, start=1):
                page_text = page.extract_text() or ""
                for match in pattern.finditer(page_text):
                    start = max(match.start() - context_chars, 0)
                    end = min(match.end() + context_chars, len(page_text))
                    results.append(
                        {
                            "page_number": page_index,
                            "match": match.group(0),
                            "context": page_text[start:end].replace("\n", " "),
                        }
                    )
    except Exception as exc:
        log_message(f"Failed to search PDF {pdf_path}: {exc}", level="ERROR")
        raise
    return results


_MARKDOWN_PROTECTED = re.compile(
    r"```[\s\S]*?```|`[^`\n]*`|\]\([^)]*\)|https?://[^\s)>]+",
    re.IGNORECASE,
)
_CANDIDATE_COLUMNS = CANDIDATE_COLUMNS


def _require_openpyxl() -> None:
    if openpyxl is None:
        raise RuntimeError(
            "openpyxl is required for XLSX operations. "
            "Install the dependencies from requirements.txt."
        )


def _xlsx_path(file_path: Path | str) -> Path:
    path = Path(file_path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx files are supported")
    if not path.is_file():
        raise FileNotFoundError(path)
    return path


def _default_output_path(input_path: Path, suffix: str, extension: str) -> Path:
    return input_path.with_name(f"{input_path.stem}_{suffix}{extension}")


def _check_output_path(output_path: Path, input_path: Path, overwrite: bool) -> None:
    if output_path.resolve() == input_path.resolve():
        raise ValueError("The output path must differ from the source path")
    if output_path.exists() and not overwrite:
        raise FileExistsError(output_path)


def _candidate_rows(candidates) -> List[Dict[str, Any]]:
    return [
        {
            "original_value": candidate.original_value,
            "replacement_value": candidate.replacement_value,
            "value_type": candidate.value_type,
            "worksheet": "; ".join(candidate.source_documents),
            "cell": "; ".join(candidate.locations),
            "occurrences": candidate.occurrences,
        }
        for candidate in candidates
    ]


def identify_xlsx_strings(
    file_path: Path | str,
    output_path: Optional[Path | str] = None,
    *,
    overwrite: bool = False,
    countries: Optional[Sequence[str]] = None,
    use_name_datasets: Optional[bool] = None,
    cache_dir: Optional[Path | str] = None,
    offline: Optional[bool] = None,
    debug: bool = False,
    name_exclusions: Optional[Sequence[str]] = None,
) -> Path:
    """Identify string cell values and write a reviewed-candidate CSV."""
    _require_openpyxl()
    input_path = _xlsx_path(file_path)
    candidate_path = Path(output_path) if output_path is not None else _default_output_path(
        input_path, "candidates", ".csv"
    )
    _check_output_path(candidate_path, input_path, overwrite)

    name_catalog = load_name_catalog(
        countries,
        use_name_datasets=use_name_datasets,
        cache_dir=cache_dir,
        offline=offline,
        debug=debug,
        name_exclusions=name_exclusions,
    )
    collector = CandidateCollector(name_catalog)
    workbook = openpyxl.load_workbook(input_path, data_only=False, keep_links=True)
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f" or not isinstance(cell.value, str):
                        continue
                    if not cell.value or is_date_string(cell.value):
                        continue
                    collector.add(cell.value, worksheet.title, cell.coordinate)
    finally:
        workbook.close()

    write_csv(candidate_path, _CANDIDATE_COLUMNS, _candidate_rows(collector.values()))
    return candidate_path


def explain_xlsx_replacement(
    file_path: Path | str,
    mapping_path: Path | str,
    replacement: str,
) -> List[Dict[str, Any]]:
    """List every cell whose anonymized value contains ``replacement``."""
    _require_openpyxl()
    input_path = _xlsx_path(file_path)
    mapping = mapping_path_rows(Path(mapping_path))
    related_values = related_mapping_values(mapping)
    results: List[Dict[str, Any]] = []

    workbook = openpyxl.load_workbook(input_path, data_only=False, read_only=True)
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if (
                        cell.data_type == "f"
                        or not isinstance(value, str)
                        or is_date_string(value)
                    ):
                        continue
                    result = replace_related_values(value, mapping, related_values)
                    if replacement not in result:
                        continue
                    if mapping.get(value) == replacement:
                        keys = [value]
                    else:
                        keys = sorted(
                            key
                            for key in related_values
                            if mapping.get(key) == replacement and key in value
                        )
                    results.append(
                        {
                            "worksheet": worksheet.title,
                            "cell": cell.coordinate,
                            "original": value,
                            "result": result,
                            "matched_keys": keys,
                        }
                    )
    finally:
        workbook.close()
    return results


def _markdown_editable_parts(content: str):
    position = 0
    for match in _MARKDOWN_PROTECTED.finditer(content):
        if match.start() > position:
            yield True, content[position : match.start()]
        yield False, match.group(0)
        position = match.end()
    if position < len(content):
        yield True, content[position:]


def _markdown_protected_spans(content: str):
    return [(match.start(), match.end()) for match in _MARKDOWN_PROTECTED.finditer(content)]


def _text_candidate_rows(
    content: str,
    document_name: str,
    *,
    name_catalog=None,
    anonymize_mode: str = "custom",
    language: str = "en",
    presidio_model: str = "de_core_news_sm",
    presidio_score_threshold: float = 0.5,
    presidio_entities: Optional[Sequence[str]] = None,
) -> List[Dict[str, Any]]:
    protected_spans = _markdown_protected_spans(content)
    if anonymize_mode == "custom":
        candidates = detect_text_candidates(
            content,
            document_name,
            protected_spans,
            name_catalog=name_catalog,
        )
    elif anonymize_mode == "presidio":
        candidates = detect_presidio_candidates(
            content,
            document_name,
            protected_spans,
            language=language,
            model_name=presidio_model,
            score_threshold=presidio_score_threshold,
            entities=presidio_entities,
        )
    else:
        raise ValueError(f"Unsupported anonymize mode: {anonymize_mode!r}")
    return _candidate_rows(candidates)


def identify_text_strings(
    file_path: Path | str,
    output_path: Optional[Path | str] = None,
    *,
    overwrite: bool = False,
    countries: Optional[Sequence[str]] = None,
    use_name_datasets: Optional[bool] = None,
    cache_dir: Optional[Path | str] = None,
    offline: Optional[bool] = None,
    debug: bool = False,
    name_exclusions: Optional[Sequence[str]] = None,
    anonymize_mode: str = "custom",
    language: str = "en",
    presidio_model: str = "de_core_news_sm",
    presidio_score_threshold: float = 0.5,
    presidio_entities: Optional[Sequence[str]] = None,
) -> Path:
    """Identify text candidates while leaving markup-specific protection here."""
    input_path = Path(file_path)
    if not input_path.is_file():
        raise FileNotFoundError(input_path)
    candidate_path = Path(output_path) if output_path is not None else _default_output_path(
        input_path, "candidates", ".csv"
    )
    _check_output_path(candidate_path, input_path, overwrite)
    name_catalog = load_name_catalog(
        countries,
        use_name_datasets=use_name_datasets,
        cache_dir=cache_dir,
        offline=offline,
        debug=debug,
        name_exclusions=name_exclusions,
    )
    write_csv(
        candidate_path,
        _CANDIDATE_COLUMNS,
        _text_candidate_rows(
            input_path.read_text(encoding="utf-8"),
            "Markdown",
            name_catalog=name_catalog,
            anonymize_mode=anonymize_mode,
            language=language,
            presidio_model=presidio_model,
            presidio_score_threshold=presidio_score_threshold,
            presidio_entities=presidio_entities,
        ),
    )
    return candidate_path


def _ignore_path_for_mapping(mapping_path: Path) -> Path:
    stem = mapping_path.stem
    for suffix in ("-mapping", "_mapping"):
        if stem.endswith(suffix):
            stem = stem[: -len(suffix)]
            break
    return mapping_path.with_name(f"{stem}-ignore.csv")


def _read_ignore_rows(ignore_path: Path) -> List[Dict[str, str]]:
    if not ignore_path.exists():
        return []
    with ignore_path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None or "original_value" not in reader.fieldnames:
            raise ValueError("Ignore CSV must contain an original_value column")
        return [
            {
                "original_value": row.get("original_value", "") or "",
                "value_type": row.get("value_type", "") or "",
            }
            for row in reader
            if row.get("original_value")
        ]


def _write_ignore_rows(ignore_path: Path, rows: List[Dict[str, str]]) -> None:
    ignore_path.parent.mkdir(parents=True, exist_ok=True)
    with ignore_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["original_value", "value_type"],
        )
        writer.writeheader()
        writer.writerows(rows)


def update_text_mapping(
    file_path: Path | str,
    mapping_path: Path | str,
    *,
    countries: Optional[Sequence[str]] = None,
    use_name_datasets: Optional[bool] = None,
    cache_dir: Optional[Path | str] = None,
    offline: Optional[bool] = None,
    debug: bool = False,
    name_exclusions: Optional[Sequence[str]] = None,
    anonymize_mode: str = "custom",
    language: str = "en",
    presidio_model: str = "de_core_news_sm",
    presidio_score_threshold: float = 0.5,
    presidio_entities: Optional[Sequence[str]] = None,
) -> Path:
    """Add newly found text candidates with status ``new``."""
    input_path = Path(file_path)
    mapping_file = Path(mapping_path)
    if not input_path.is_file():
        raise FileNotFoundError(input_path)
    existing_documents = []
    if mapping_file.exists():
        from wit_pytools.anonymization import read_mapping_rows

        existing_documents = read_mapping_rows(mapping_file)

    ignore_path = _ignore_path_for_mapping(mapping_file)
    ignore_rows = _read_ignore_rows(ignore_path)
    ignored_values = {
        row["original_value"].strip()
        for row in ignore_rows
        if row["original_value"].strip()
    }
    anon_values = set()
    rows = []
    for row in existing_documents:
        originals = {
            original.strip()
            for original in row["original_value"].split(";")
            if original.strip()
        }
        if row["status"] == "keep":
            for original in originals:
                if original not in ignored_values:
                    ignore_rows.append(
                        {
                            "original_value": original,
                            "value_type": row["value_type"],
                        }
                    )
                    ignored_values.add(original)
            continue
        if row["status"] == "anon":
            anon_values.update(originals)
            row["source_documents"] = ""
            row["locations"] = ""
            row["occurrences"] = ""
        rows.append(row)

    rows = [
        row
        for row in rows
        if not (
            row["status"] == "new"
            and any(
                original in ignored_values or original in anon_values
                for original in row["original_value"].split(";")
            )
        )
    ]
    unique_ignore_rows = []
    seen_ignore_values = set()
    for row in ignore_rows:
        original = row["original_value"].strip()
        if not original or original in seen_ignore_values:
            continue
        seen_ignore_values.add(original)
        unique_ignore_rows.append(
            {"original_value": original, "value_type": row["value_type"]}
        )
    ignore_rows = unique_ignore_rows
    _write_ignore_rows(ignore_path, ignore_rows)
    known = ignored_values | anon_values | {
        original.strip()
        for row in rows
        for original in row["original_value"].split(";")
        if original.strip()
    }
    name_catalog = load_name_catalog(
        countries,
        use_name_datasets=use_name_datasets,
        cache_dir=cache_dir,
        offline=offline,
        debug=debug,
        name_exclusions=name_exclusions,
    )
    candidates = _text_candidate_rows(
        input_path.read_text(encoding="utf-8"),
        input_path.name,
        name_catalog=name_catalog,
        anonymize_mode=anonymize_mode,
        language=language,
        presidio_model=presidio_model,
        presidio_score_threshold=presidio_score_threshold,
    )
    for candidate in candidates:
        if candidate["original_value"] in known:
            continue
        rows.append(
            {
                "replacement_value": candidate["replacement_value"],
                "original_value": candidate["original_value"],
                "value_type": candidate["value_type"],
                "status": "new",
                "source_documents": candidate["worksheet"],
                "locations": candidate["cell"],
                "occurrences": str(candidate["occurrences"]),
            }
        )
        known.add(candidate["original_value"])
    mapping_file.parent.mkdir(parents=True, exist_ok=True)
    with mapping_file.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "status",
                "replacement_value",
                "original_value",
                "value_type",
                "source_documents",
                "locations",
                "occurrences",
            ],
        )
        writer.writeheader()
        writer.writerows(rows)
    return mapping_file


def mapping_matches_text(content: str, mapping_path: Path | str) -> bool:
    """Return whether approved mappings match editable Markdown/text."""
    mapping = mapping_path_rows(Path(mapping_path))
    return mapping_matches_parts(_markdown_editable_parts(content), mapping)


def anonymize_text_content(content: str, mapping_path: Path | str) -> str:
    """Apply a mapping to editable Markdown/text content."""
    mapping = mapping_path_rows(Path(mapping_path))
    return replace_text_parts(_markdown_editable_parts(content), mapping)


def anonymize_text(
    file_path: Path | str,
    mapping_path: Path | str,
    output_path: Optional[Path | str] = None,
    *,
    overwrite: bool = False,
) -> Path:
    """Apply a mapping to Markdown/text and write an output file."""
    input_path = Path(file_path)
    if not input_path.is_file():
        raise FileNotFoundError(input_path)
    output = Path(output_path) if output_path is not None else input_path.with_name(
        f"{input_path.stem}_anon{input_path.suffix}"
    )
    _check_output_path(output, input_path, overwrite)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        anonymize_text_content(input_path.read_text(encoding="utf-8"), mapping_path),
        encoding="utf-8",
    )
    return output


def _anonymized_path_for_mapping(mapping_path: Path) -> Path:
    suffix = "_mapping"
    stem = mapping_path.stem
    if stem.endswith(suffix):
        stem = stem[: -len(suffix)]
    return mapping_path.with_name(f"{stem}_anonymized.xlsx")


def anonymize_xlsx(
    file_path: Path | str,
    mapping_path: Path | str,
    output_path: Optional[Path | str] = None,
    *,
    overwrite: bool = False,
) -> Path:
    """Apply a reviewed mapping to string cells in an XLSX workbook."""
    _require_openpyxl()
    input_path = _xlsx_path(file_path)
    reviewed_mapping_path = Path(mapping_path)
    if not reviewed_mapping_path.is_file():
        raise FileNotFoundError(reviewed_mapping_path)
    mapping = mapping_path_rows(reviewed_mapping_path)
    anonymized_path = (
        Path(output_path)
        if output_path is not None
        else _anonymized_path_for_mapping(reviewed_mapping_path)
    )
    if anonymized_path.suffix.lower() != ".xlsx":
        raise ValueError("The anonymized output must be an .xlsx file")
    _check_output_path(anonymized_path, input_path, overwrite)

    workbook = openpyxl.load_workbook(input_path, data_only=False, keep_links=True)
    related_values = related_mapping_values(mapping)
    temporary_path = None
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if (
                        cell.data_type != "f"
                        and isinstance(cell.value, str)
                        and not is_date_string(cell.value)
                    ):
                        replacement = replace_related_values(
                            cell.value, mapping, related_values
                        )
                        if replacement != cell.value:
                            cell.value = replacement

        anonymized_path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            dir=anonymized_path.parent,
            prefix=f".{anonymized_path.stem}-",
            suffix=".xlsx",
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, anonymized_path)
    finally:
        workbook.close()
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()
    return anonymized_path
